from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Optional
import os
import secrets
import shutil
import time

from fastapi import (
    FastAPI,
    Depends,
    HTTPException,
    UploadFile,
    File,
    WebSocket,
    WebSocketDisconnect,
)
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from sqlalchemy import (
    create_engine,
    String,
    Integer,
    Boolean,
    DateTime,
    ForeignKey,
    UniqueConstraint,
    select,
    func,
    delete,
    text,
)
from sqlalchemy.orm import (
    DeclarativeBase,
    Mapped,
    mapped_column,
    sessionmaker,
    Session,
)
from passlib.context import CryptContext
from openpyxl import load_workbook


# ============================================================
# CONFIGURATION
# ============================================================

BASE = Path(__file__).resolve().parent.parent
UPLOADS = BASE / "uploads" / "questions"
UPLOADS.mkdir(parents=True, exist_ok=True)

DB_URL = os.getenv(
    "DATABASE_URL",
    f"sqlite:///{BASE / 'funfriday.db'}",
)

engine = create_engine(
    DB_URL,
    connect_args=(
        {"check_same_thread": False}
        if DB_URL.startswith("sqlite")
        else {}
    ),
)

SessionLocal = sessionmaker(
    bind=engine,
    autoflush=False,
    expire_on_commit=False,
)

pwd = CryptContext(schemes=["bcrypt"], deprecated="auto")

# Demo/local token store.
# Use a secure session/JWT implementation for production.
TOKENS: dict[str, int] = {}
PRESENCE: dict[int, float] = {}

MAX_IMAGE_BYTES = 5 * 1024 * 1024


# ============================================================
# DATABASE MODELS
# ============================================================

class Base(DeclarativeBase):
    pass


class User(Base):
    __tablename__ = "users"

    id: Mapped[int] = mapped_column(primary_key=True)
    username: Mapped[str] = mapped_column(
        String(80), unique=True, index=True
    )
    password_hash: Mapped[str] = mapped_column(String(255))
    role: Mapped[str] = mapped_column(String(20), default="user")
    is_active: Mapped[bool] = mapped_column(Boolean, default=True)


class Team(Base):
    __tablename__ = "teams"

    id: Mapped[int] = mapped_column(primary_key=True)
    name: Mapped[str] = mapped_column(String(100))
    created_at: Mapped[datetime] = mapped_column(
        DateTime, default=datetime.utcnow
    )


class Membership(Base):
    __tablename__ = "memberships"

    id: Mapped[int] = mapped_column(primary_key=True)
    user_id: Mapped[int] = mapped_column(
        ForeignKey("users.id"), unique=True
    )
    team_id: Mapped[int] = mapped_column(ForeignKey("teams.id"))


class Activity(Base):
    __tablename__ = "activities"

    id: Mapped[int] = mapped_column(primary_key=True)
    key: Mapped[str] = mapped_column(String(40), unique=True)
    name: Mapped[str] = mapped_column(String(100))
    duration_seconds: Mapped[int] = mapped_column(
        Integer, default=60
    )
    unlocked: Mapped[bool] = mapped_column(Boolean, default=False)
    published: Mapped[bool] = mapped_column(Boolean, default=False)
    input_type: Mapped[Optional[str]] = mapped_column(
        String(20), nullable=True, default="image"
    )
    question_mode: Mapped[Optional[str]] = mapped_column(
        String(20), nullable=True, default="team"
    )
    question_type: Mapped[Optional[str]] = mapped_column(
        String(30), nullable=True, default="buzzer"
    )


class Question(Base):
    __tablename__ = "questions"

    id: Mapped[int] = mapped_column(primary_key=True)
    number: Mapped[int] = mapped_column(Integer)
    activity_key: Mapped[Optional[str]] = mapped_column(
        String(40), default="find-object", nullable=True
    )
    prompt: Mapped[Optional[str]] = mapped_column(
        String(500), nullable=True
    )
    media_url: Mapped[Optional[str]] = mapped_column(
        String(500), nullable=True
    )
    target_user_id: Mapped[Optional[int]] = mapped_column(
        ForeignKey("users.id"), nullable=True
    )
    target_team_id: Mapped[Optional[int]] = mapped_column(
        ForeignKey("teams.id"), nullable=True
    )
    answer: Mapped[str] = mapped_column(String(250))
    hint: Mapped[Optional[str]] = mapped_column(
        String(250), nullable=True
    )
    status: Mapped[str] = mapped_column(String(20), default="locked")
    started_at: Mapped[Optional[datetime]] = mapped_column(
        DateTime, nullable=True
    )
    ends_at: Mapped[Optional[datetime]] = mapped_column(
        DateTime, nullable=True
    )
    resolved: Mapped[bool] = mapped_column(Boolean, default=False)


class QImage(Base):
    __tablename__ = "question_images"

    id: Mapped[int] = mapped_column(primary_key=True)
    question_id: Mapped[int] = mapped_column(
        ForeignKey("questions.id")
    )
    path: Mapped[str] = mapped_column(String(500))
    order_num: Mapped[int] = mapped_column(Integer, default=0)


class Submission(Base):
    __tablename__ = "submissions"

    id: Mapped[int] = mapped_column(primary_key=True)
    question_id: Mapped[int] = mapped_column(
        ForeignKey("questions.id")
    )
    team_id: Mapped[int] = mapped_column(ForeignKey("teams.id"))
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id"))
    answer_text: Mapped[str] = mapped_column(String(500))
    status: Mapped[str] = mapped_column(String(20), default="pending")
    submitted_at: Mapped[datetime] = mapped_column(
        DateTime, default=datetime.utcnow
    )
    points: Mapped[int] = mapped_column(Integer, default=0)


class Topic(Base):
    __tablename__ = "topics"

    id: Mapped[int] = mapped_column(primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id"))
    text: Mapped[str] = mapped_column(String(250))
    status: Mapped[str] = mapped_column(String(20), default="locked")
    started_at: Mapped[Optional[datetime]] = mapped_column(
        DateTime, nullable=True
    )
    ends_at: Mapped[Optional[datetime]] = mapped_column(
        DateTime, nullable=True
    )
    resolved: Mapped[bool] = mapped_column(Boolean, default=False)


class Buzzer(Base):
    __tablename__ = "buzzer_entries"
    __table_args__ = (
        UniqueConstraint(
            "round_key",
            "team_id",
            name="uq_round_team_buzzer",
        ),
    )

    id: Mapped[int] = mapped_column(primary_key=True)
    round_key: Mapped[str] = mapped_column(String(100), index=True)
    activity_key: Mapped[str] = mapped_column(String(40))
    team_id: Mapped[int] = mapped_column(ForeignKey("teams.id"))
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id"))
    pressed_at: Mapped[datetime] = mapped_column(
        DateTime, default=datetime.utcnow
    )
    position: Mapped[int] = mapped_column(Integer)


class Score(Base):
    __tablename__ = "scores"

    id: Mapped[int] = mapped_column(primary_key=True)
    team_id: Mapped[int] = mapped_column(ForeignKey("teams.id"))
    activity_key: Mapped[str] = mapped_column(String(40))
    round_key: Mapped[str] = mapped_column(String(100))
    points: Mapped[int] = mapped_column(Integer, default=0)
    updated_by: Mapped[int] = mapped_column(ForeignKey("users.id"))


class ImportedResult(Base):
    __tablename__ = "imported_results"

    id: Mapped[int] = mapped_column(primary_key=True)
    team_name: Mapped[str] = mapped_column(
        String(100), unique=True
    )
    object_score: Mapped[int] = mapped_column(Integer, default=0)
    speech_score: Mapped[int] = mapped_column(Integer, default=0)
    total: Mapped[int] = mapped_column(Integer, default=0)


Base.metadata.create_all(engine)


def run_sqlite_migrations():
    if not DB_URL.startswith("sqlite"):
        return
    try:
        with engine.connect() as conn:
            res = conn.execute(text("PRAGMA table_info(questions)")).fetchall()
            col_names = [r[1] for r in res]
            if "activity_key" not in col_names:
                conn.execute(text("ALTER TABLE questions ADD COLUMN activity_key VARCHAR(40) DEFAULT 'find-object'"))
            if "prompt" not in col_names:
                conn.execute(text("ALTER TABLE questions ADD COLUMN prompt VARCHAR(500)"))
            if "media_url" not in col_names:
                conn.execute(text("ALTER TABLE questions ADD COLUMN media_url VARCHAR(500)"))
            if "target_user_id" not in col_names:
                conn.execute(text("ALTER TABLE questions ADD COLUMN target_user_id INTEGER"))
            if "target_team_id" not in col_names:
                conn.execute(text("ALTER TABLE questions ADD COLUMN target_team_id INTEGER"))

            act_res = conn.execute(text("PRAGMA table_info(activities)")).fetchall()
            act_cols = [r[1] for r in act_res]
            if "question_type" not in act_cols:
                conn.execute(text("ALTER TABLE activities ADD COLUMN question_type VARCHAR(30) DEFAULT 'buzzer'"))
            conn.commit()
    except Exception as e:
        print("Migration error:", e)

run_sqlite_migrations()


# ============================================================
# FASTAPI APPLICATION
# ============================================================

app = FastAPI(
    title="Fun Friday Arena API",
    version="1.0.0",
)

app.mount(
    "/uploads",
    StaticFiles(directory=str(BASE / "uploads")),
    name="uploads",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=os.getenv(
        "CORS_ORIGINS",
        "http://localhost:5173,http://127.0.0.1:5173",
    ).split(","),
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

bearer = HTTPBearer(auto_error=False)


# ============================================================
# DATABASE AND AUTHENTICATION HELPERS
# ============================================================

def db_dep():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def now():
    return datetime.utcnow()


def to_iso(dt: Optional[datetime]) -> Optional[str]:
    if not dt:
        return None
    if dt.tzinfo is None:
        return dt.isoformat() + "Z"
    return dt.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.%f")[:-3] + "Z"


def is_online(uid: int):
    return time.time() - PRESENCE.get(uid, 0) < 35


def user_dict(u: User, db: Session):
    if u.role == "admin":
        return {
            "id": u.id,
            "username": u.username,
            "role": u.role,
            "online": is_online(u.id),
            "team_id": None,
            "team_name": None,
        }

    membership = db.scalar(
        select(Membership).where(Membership.user_id == u.id)
    )
    team = db.get(Team, membership.team_id) if membership else None

    return {
        "id": u.id,
        "username": u.username,
        "role": u.role,
        "online": is_online(u.id),
        "team_id": team.id if team else None,
        "team_name": team.name if team else None,
    }


def current_user(
    creds: HTTPAuthorizationCredentials = Depends(bearer),
    db: Session = Depends(db_dep),
):
    if not creds or creds.credentials not in TOKENS:
        raise HTTPException(
            status_code=401,
            detail="Please log in",
        )

    user = db.get(User, TOKENS[creds.credentials])

    if not user or not user.is_active:
        raise HTTPException(
            status_code=401,
            detail="Account unavailable",
        )

    PRESENCE[user.id] = time.time()
    return user


def admin(u: User = Depends(current_user)):
    if u.role != "admin":
        raise HTTPException(
            status_code=403,
            detail="Admin access required",
        )
    return u


def team_for(db: Session, uid: int):
    user = db.get(User, uid)
    if user and user.role == "admin":
        return None
    membership = db.scalar(
        select(Membership).where(Membership.user_id == uid)
    )
    return db.get(Team, membership.team_id) if membership else None


def active_round(db: Session, key: str):
    if key.startswith("object:"):
        question = db.get(Question, int(key.split(":")[1]))
        return question, "find-object"

    if key.startswith("speech:"):
        topic = db.get(Topic, int(key.split(":")[1]))
        return topic, "lightning-speech"

    return None, None


def round_open(round_item):
    if (
        not round_item
        or round_item.status != "active"
        or round_item.resolved
    ):
        return False

    if round_item.ends_at and now() >= round_item.ends_at:
        round_item.status = "closed"
        return False

    return True


# ============================================================
# REQUEST SCHEMAS
# ============================================================

class LoginIn(BaseModel):
    username: str
    password: str


class TimerIn(BaseModel):
    duration_seconds: int = Field(ge=10, le=3600)


class TeamGenerateIn(BaseModel):
    user_ids: Optional[list[int]] = None
    team_size: int = Field(default=2, ge=1, le=50)
    include_online: bool = True
    include_offline: bool = True


class RenameIn(BaseModel):
    name: str = Field(min_length=1, max_length=100)


class QuestionIn(BaseModel):
    number: int = Field(ge=1)
    answer: str = Field(min_length=1, max_length=250)
    hint: Optional[str] = None
    activity_key: Optional[str] = "find-object"
    prompt: Optional[str] = None
    media_url: Optional[str] = None
    target_user_id: Optional[int] = None
    target_team_id: Optional[int] = None


class AnswerIn(BaseModel):
    answer_text: str = Field(min_length=1, max_length=500)


class EvaluateIn(BaseModel):
    status: str
    points: int = Field(default=0, ge=-100000, le=100000)


class TopicIn(BaseModel):
    user_id: int
    text: str = Field(min_length=1, max_length=250)


class ScoreIn(BaseModel):
    team_id: int
    activity_key: Optional[str] = "find-object"
    round_key: str
    points: int = Field(default=0, ge=-100000, le=100000)


class PublishIn(BaseModel):
    published: bool


class NewUserIn(BaseModel):
    username: str = Field(
        min_length=2,
        max_length=80,
        pattern=r"^[A-Za-z0-9_.-]+$",
    )
    password: str = Field(min_length=8, max_length=128)


class TeamCreateIn(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    member_ids: list[int] = []


class BuzzerIn(BaseModel):
    round_key: str
    activity_key: Optional[str] = "find-object"


class AnswerSubmitIn(BaseModel):
    question_id: int
    answer_text: str = Field(min_length=1, max_length=500)


class ActivityUnlockIn(BaseModel):
    published: Optional[bool] = None
    unlocked: Optional[bool] = None


class ActivityCreateIn(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    key: Optional[str] = None
    duration_seconds: int = Field(default=60, ge=10, le=3600)
    input_type: str = Field(default="image")
    question_mode: str = Field(default="team")
    question_type: Optional[str] = "buzzer"


class ActivityRenameIn(BaseModel):
    name: str = Field(min_length=1, max_length=100)


# ============================================================
# STARTUP
# ============================================================

@app.on_event("startup")
def seed():
    with SessionLocal() as db:
        activities = [
            ("find-object", "Find the Object", 60),
            ("lightning-speech", "Lightning Speech", 120),
        ]

        for key, name, duration in activities:
            if not db.scalar(
                select(Activity).where(Activity.key == key)
            ):
                db.add(
                    Activity(
                        key=key,
                        name=name,
                        duration_seconds=duration,
                    )
                )

        if not db.scalar(
            select(User).where(User.username == "admin")
        ):
            admin_password = os.getenv(
                "ADMIN_PASSWORD",
                "Admin123!",
            )

            db.add(
                User(
                    username="admin",
                    password_hash=pwd.hash(admin_password),
                    role="admin",
                )
            )

            if os.getenv("SEED_DEMO_USERS", "true").lower() != "false":
                user_password = os.getenv(
                    "DEMO_USER_PASSWORD",
                    "User123!",
                )

                db.add_all(
                    [
                        User(
                            username=username,
                            password_hash=pwd.hash(user_password),
                        )
                        for username in (
                            "alice",
                            "bob",
                            "charlie",
                            "dina",
                        )
                    ]
                )

        db.commit()


# ============================================================
# HEALTH AND AUTHENTICATION
# ============================================================

@app.get("/api/health")
def health():
    return {"ok": True, "app": "Fun Friday Arena"}


@app.post("/api/auth/signup")
def signup(
    data: NewUserIn,
    db: Session = Depends(db_dep),
):
    name = data.username.strip()

    if db.scalar(select(User).where(User.username == name)):
        raise HTTPException(
            status_code=409,
            detail="Username already exists. Choose another username.",
        )

    user = User(
        username=name,
        password_hash=pwd.hash(data.password),
        role="user",
        is_active=True,
    )

    db.add(user)

    try:
        db.commit()
    except Exception:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Username already exists. Choose another username.",
        )

    db.refresh(user)

    return {
        "ok": True,
        "message": "Account created. You can now log in.",
    }


@app.post("/api/auth/login")
def login(
    data: LoginIn,
    db: Session = Depends(db_dep),
):
    user = db.scalar(
        select(User).where(User.username == data.username)
    )

    if (
        not user
        or not user.is_active
        or not pwd.verify(data.password, user.password_hash)
    ):
        raise HTTPException(
            status_code=401,
            detail="Invalid username or password",
        )

    token = secrets.token_urlsafe(32)
    TOKENS[token] = user.id
    PRESENCE[user.id] = time.time()

    return {
        "token": token,
        "user": user_dict(user, db),
    }


@app.post("/api/auth/logout")
def logout(
    u: User = Depends(current_user),
    creds: HTTPAuthorizationCredentials = Depends(bearer),
):
    TOKENS.pop(creds.credentials, None)
    PRESENCE.pop(u.id, None)

    return {"ok": True}


@app.get("/api/me")
def me(
    u: User = Depends(current_user),
    db: Session = Depends(db_dep),
):
    return user_dict(u, db)


@app.post("/api/heartbeat")
def heartbeat(u: User = Depends(current_user)):
    PRESENCE[u.id] = time.time()
    return {"ok": True}


# ============================================================
# USER MANAGEMENT
# ============================================================

@app.get("/api/users")
def users(
    u: User = Depends(current_user),
    db: Session = Depends(db_dep),
):
    """Return active users to authenticated callers."""

    active_users = db.scalars(
        select(User)
        .where(User.is_active == True)
        .order_by(User.username)
    ).all()

    return [user_dict(user, db) for user in active_users]


@app.post("/api/users")
def create_user(
    data: NewUserIn,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    """Create a participant from the admin interface."""

    name = data.username.strip()

    if db.scalar(select(User).where(User.username == name)):
        raise HTTPException(
            status_code=409,
            detail="Username already exists",
        )

    user = User(
        username=name,
        password_hash=pwd.hash(data.password),
        role="user",
        is_active=True,
    )

    db.add(user)

    try:
        db.commit()
        db.refresh(user)
    except Exception:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Could not create user. The username may already exist.",
        )

    return user_dict(user, db)


@app.delete("/api/users/{user_id}")
def delete_user(
    user_id: int,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    """
    Soft-delete a participant:
    - Mark the account inactive.
    - Remove team membership.
    - Invalidate existing login tokens.
    """

    target = db.get(User, user_id)

    if not target or not target.is_active:
        raise HTTPException(
            status_code=404,
            detail="User not found",
        )

    if target.role == "admin":
        raise HTTPException(
            status_code=403,
            detail="Cannot delete an admin account",
        )

    # Mark the account inactive instead of permanently deleting it.
    target.is_active = False

    # Remove the participant's team membership.
    db.query(Membership).filter(
        Membership.user_id == target.id
    ).delete(synchronize_session=False)

    try:
        db.commit()
    except Exception:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail="Could not delete the participant",
        )

    # Invalidate the deleted user's existing tokens.
    tokens_to_remove = [
        token
        for token, uid in TOKENS.items()
        if uid == target.id
    ]

    for token in tokens_to_remove:
        TOKENS.pop(token, None)

    PRESENCE.pop(target.id, None)

    return {
        "ok": True,
        "message": f"User {target.username} deleted successfully",
        "user_id": target.id,
    }


# ============================================================
# TEAM MANAGEMENT
# ============================================================

@app.post("/api/teams/generate")
def generate_teams(
    data: TeamGenerateIn,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    chosen = []

    if data.user_ids and len(data.user_ids) > 0:
        candidate_ids = dict.fromkeys(data.user_ids)
    else:
        candidate_ids = [
            u.id
            for u in db.scalars(
                select(User).where(User.is_active == True, User.role == "user")
            ).all()
        ]

    for uid in candidate_ids:
        user = db.get(User, uid)

        if not user or not user.is_active or user.role != "user":
            continue

        online = is_online(uid)

        if (
            (online and data.include_online)
            or (not online and data.include_offline)
        ):
            chosen.append(user)

    if not chosen:
        raise HTTPException(
            status_code=400,
            detail="Select at least one eligible participant",
        )

    import random
    random.shuffle(chosen)

    db.query(Membership).delete()
    db.query(Team).delete()
    db.flush()

    s = max(1, data.team_size)
    n = len(chosen)

    if n <= s:
        groups = [chosen]
    else:
        num_teams = n // s
        remainder = n % s
        groups = []
        cur_idx = 0
        for i in range(num_teams):
            extra = 1 if i < remainder else 0
            size = s + extra
            groups.append(chosen[cur_idx : cur_idx + size])
            cur_idx += size

    output = []

    for index, group in enumerate(groups, 1):
        team = Team(name=f"Team {index}")
        db.add(team)
        db.flush()

        for user in group:
            db.add(
                Membership(
                    team_id=team.id,
                    user_id=user.id,
                )
            )

        output.append(
            {
                "id": team.id,
                "name": team.name,
                "members": [user.username for user in group],
            }
        )

    db.commit()

    return {
        "teams": output,
        "count": len(chosen),
    }


@app.get("/api/teams")
def get_teams(
    u: User = Depends(current_user),
    db: Session = Depends(db_dep),
):
    result = []

    for team in db.scalars(
        select(Team).where(Team.name != "Admin Staff").order_by(Team.id)
    ).all():
        members = []

        for membership in db.scalars(
            select(Membership).where(
                Membership.team_id == team.id
            )
        ).all():
            user = db.get(User, membership.user_id)

            if user and user.is_active and user.role != "admin":
                members.append(
                    {
                        "id": user.id,
                        "username": user.username,
                    }
                )

        result.append(
            {
                "id": team.id,
                "name": team.name,
                "members": members,
            }
        )

    return result


@app.post("/api/teams")
def create_team(
    data: TeamCreateIn,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    name = data.name.strip()
    team = Team(name=name)
    db.add(team)
    db.flush()

    for user_id in data.member_ids:
        user = db.get(User, user_id)
        if user and user.is_active and user.role != "admin":
            db.query(Membership).filter(Membership.user_id == user.id).delete()
            db.add(Membership(team_id=team.id, user_id=user.id))

    db.commit()
    db.refresh(team)

    members = [
        {"id": m.user_id, "username": db.get(User, m.user_id).username}
        for m in db.scalars(
            select(Membership).where(Membership.team_id == team.id)
        ).all()
        if db.get(User, m.user_id)
    ]

    return {
        "id": team.id,
        "name": team.name,
        "members": members,
    }


@app.patch("/api/teams/{team_id}")
def rename_team(
    team_id: int,
    data: RenameIn,
    u: User = Depends(current_user),
    db: Session = Depends(db_dep),
):
    team = db.get(Team, team_id)
    if not team:
        raise HTTPException(status_code=404, detail="Team not found")

    is_admin = (u.role == "admin")
    is_member = db.scalar(
        select(Membership).where(
            Membership.team_id == team_id,
            Membership.user_id == u.id,
        )
    ) is not None

    if not is_admin and not is_member:
        raise HTTPException(
            status_code=403,
            detail="You must be an admin or a member of this team to rename it",
        )

    team.name = data.name.strip()
    db.commit()
    return {"ok": True, "id": team.id, "name": team.name}


@app.delete("/api/teams/{team_id}")
def delete_team(
    team_id: int,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    """Delete a team and all its memberships, scores, and buzzer entries."""
    team = db.get(Team, team_id)
    if not team:
        raise HTTPException(status_code=404, detail="Team not found")

    # Remove memberships
    db.query(Membership).filter(
        Membership.team_id == team_id
    ).delete(synchronize_session=False)

    # Remove scores for this team
    db.query(Score).filter(
        Score.team_id == team_id
    ).delete(synchronize_session=False)

    # Remove buzzer entries for this team
    db.query(Buzzer).filter(
        Buzzer.team_id == team_id
    ).delete(synchronize_session=False)

    # Remove submissions for this team
    db.query(Submission).filter(
        Submission.team_id == team_id
    ).delete(synchronize_session=False)

    # Clear any question assignments to this team
    db.query(Question).filter(
        Question.target_team_id == team_id
    ).update({"target_team_id": None}, synchronize_session=False)

    db.delete(team)

    try:
        db.commit()
    except Exception:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail="Could not delete the team",
        )

    return {
        "ok": True,
        "message": f"Team '{team.name}' deleted successfully",
        "team_id": team_id,
    }


# ============================================================
# ACTIVITY MANAGEMENT
# ============================================================

@app.get("/api/activities")
def get_activities(
    u: User = Depends(current_user),
    db: Session = Depends(db_dep),
):
    acts = db.scalars(select(Activity).order_by(Activity.id)).all()
    return [
        {
            "id": a.id,
            "key": a.key,
            "name": a.name,
            "duration_seconds": a.duration_seconds,
            "unlocked": a.unlocked,
            "published": a.published,
            "input_type": a.input_type or "image",
            "question_mode": a.question_mode or "team",
            "question_type": a.question_type or "buzzer",
        }
        for a in acts
    ]


@app.post("/api/activities")
def create_activity(
    data: ActivityCreateIn,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    """Create a new activity from the admin interface."""
    import re
    key = data.key or re.sub(r'[^a-z0-9]+', '-', data.name.strip().lower()).strip('-')

    if db.scalar(select(Activity).where(Activity.key == key)):
        raise HTTPException(
            status_code=409,
            detail=f"Activity with key '{key}' already exists",
        )

    if data.input_type not in ("image", "audio", "text"):
        raise HTTPException(
            status_code=400,
            detail="input_type must be 'image', 'audio', or 'text'",
        )

    if data.question_mode not in ("team", "member"):
        raise HTTPException(
            status_code=400,
            detail="question_mode must be 'team' or 'member'",
        )

    activity = Activity(
        key=key,
        name=data.name.strip(),
        duration_seconds=data.duration_seconds,
        input_type=data.input_type,
        question_mode=data.question_mode,
        question_type=data.question_type or "buzzer",
    )
    db.add(activity)
    db.commit()
    db.refresh(activity)

    return {
        "id": activity.id,
        "key": activity.key,
        "name": activity.name,
        "duration_seconds": activity.duration_seconds,
        "unlocked": activity.unlocked,
        "published": activity.published,
        "input_type": activity.input_type,
        "question_mode": activity.question_mode,
        "question_type": activity.question_type,
    }


@app.delete("/api/activities/{key}")
def delete_activity(
    key: str,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    """Delete an activity and all related questions, scores, and buzzers."""
    if key in ("find-object", "lightning-speech"):
        raise HTTPException(
            status_code=400,
            detail="Cannot delete built-in activities",
        )

    act = db.scalar(select(Activity).where(Activity.key == key))
    if not act:
        raise HTTPException(status_code=404, detail="Activity not found")

    # Delete all questions for this activity and their clue files
    questions = db.scalars(select(Question).where(Question.activity_key == key)).all()
    for q in questions:
        images = db.scalars(select(QImage).where(QImage.question_id == q.id)).all()
        for img in images:
            if img.path:
                file_name = Path(img.path).name
                file_path = UPLOADS / file_name
                if file_path.is_file():
                    try:
                        file_path.unlink()
                    except Exception:
                        pass
            db.delete(img)
        db.query(Submission).filter(Submission.question_id == q.id).delete(synchronize_session=False)
        db.delete(q)

    # Remove scores for this activity
    db.query(Score).filter(
        Score.activity_key == key
    ).delete(synchronize_session=False)

    # Remove buzzer entries for this activity
    db.query(Buzzer).filter(
        Buzzer.activity_key == key
    ).delete(synchronize_session=False)

    db.delete(act)
    db.commit()

    return {
        "ok": True,
        "message": f"Activity '{act.name}' deleted",
        "key": key,
    }


@app.patch("/api/activities/{key}/rename")
def rename_activity(
    key: str,
    data: ActivityRenameIn,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    """Rename an activity."""
    act = db.scalar(select(Activity).where(Activity.key == key))
    if not act:
        raise HTTPException(status_code=404, detail="Activity not found")

    act.name = data.name.strip()
    db.commit()
    return {"ok": True, "key": act.key, "name": act.name}


@app.patch("/api/activities/{key}/timer")
def update_activity_timer(
    key: str,
    data: TimerIn,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    act = db.scalar(select(Activity).where(Activity.key == key))
    if not act:
        raise HTTPException(status_code=404, detail="Activity not found")

    act.duration_seconds = data.duration_seconds
    db.commit()
    return {"ok": True, "key": act.key, "duration_seconds": act.duration_seconds}


@app.patch("/api/activities/{key}/unlock")
def unlock_activity(
    key: str,
    data: Optional[ActivityUnlockIn] = None,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    act = db.scalar(select(Activity).where(Activity.key == key))
    if not act:
        raise HTTPException(status_code=404, detail="Activity not found")

    if data and data.published is not None:
        act.unlocked = data.published
    elif data and data.unlocked is not None:
        act.unlocked = data.unlocked
    else:
        act.unlocked = not act.unlocked

    db.commit()
    return {"ok": True, "key": act.key, "unlocked": act.unlocked}


@app.patch("/api/activities/{key}/publish")
def publish_activity(
    key: str,
    data: PublishIn,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    act = db.scalar(select(Activity).where(Activity.key == key))
    if not act:
        raise HTTPException(status_code=404, detail="Activity not found")

    act.published = data.published
    db.commit()
    return {"ok": True, "key": act.key, "published": act.published}


# ============================================================
# QUESTION MANAGEMENT (FIND THE OBJECT)
# ============================================================

@app.get("/api/questions")
def get_questions(
    activity_key: Optional[str] = None,
    u: User = Depends(current_user),
    db: Session = Depends(db_dep),
):
    query = select(Question).order_by(Question.number.asc())
    if activity_key and activity_key != "all":
        query = query.where(Question.activity_key == activity_key)

    questions = db.scalars(query).all()
    output = []

    for q in questions:
        images = [
            img.path
            for img in db.scalars(
                select(QImage)
                .where(QImage.question_id == q.id)
                .order_by(QImage.order_num.asc())
            ).all()
        ]

        target_user = db.get(User, q.target_user_id) if q.target_user_id else None
        target_team = db.get(Team, q.target_team_id) if q.target_team_id else None

        output.append(
            {
                "id": q.id,
                "number": q.number,
                "activity_key": q.activity_key or "find-object",
                "prompt": q.prompt,
                "media_url": q.media_url,
                "target_user_id": q.target_user_id,
                "target_team_id": q.target_team_id,
                "target_username": target_user.username if target_user else None,
                "target_team_name": target_team.name if target_team else None,
                "answer": q.answer if u.role == "admin" else None,
                "hint": q.hint,
                "status": q.status,
                "started_at": to_iso(q.started_at),
                "ends_at": to_iso(q.ends_at),
                "resolved": q.resolved,
                "images": images,
            }
        )

    return output


@app.post("/api/questions")
def create_question(
    data: QuestionIn,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    act_key = data.activity_key or "find-object"
    q = Question(
        number=data.number,
        activity_key=act_key,
        prompt=data.prompt.strip() if data.prompt else None,
        media_url=data.media_url.strip() if data.media_url else None,
        target_user_id=data.target_user_id,
        target_team_id=data.target_team_id,
        answer=data.answer.strip(),
        hint=data.hint.strip() if data.hint else None,
        status="locked",
    )
    db.add(q)
    db.commit()
    db.refresh(q)

    target_user = db.get(User, q.target_user_id) if q.target_user_id else None
    target_team = db.get(Team, q.target_team_id) if q.target_team_id else None

    return {
        "id": q.id,
        "number": q.number,
        "activity_key": q.activity_key,
        "prompt": q.prompt,
        "media_url": q.media_url,
        "target_user_id": q.target_user_id,
        "target_team_id": q.target_team_id,
        "target_username": target_user.username if target_user else None,
        "target_team_name": target_team.name if target_team else None,
        "answer": q.answer,
        "hint": q.hint,
        "status": q.status,
        "images": [],
    }


@app.post("/api/questions/{question_id}/images")
async def upload_question_images(
    question_id: int,
    files: list[UploadFile] = File(default=[]),
    file: Optional[UploadFile] = File(default=None),
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    question = db.get(Question, question_id)
    if not question:
        raise HTTPException(status_code=404, detail="Question not found")

    all_files = list(files)
    if file:
        all_files.append(file)

    if not all_files:
        raise HTTPException(status_code=400, detail="No media files provided")

    saved_images = []
    existing_max = db.scalar(
        select(func.max(QImage.order_num)).where(QImage.question_id == question_id)
    )
    start_order = (existing_max + 1) if existing_max is not None else 0

    for idx, f in enumerate(all_files):
        ext = Path(f.filename or "media.png").suffix or ".png"
        filename = f"q_{question_id}_{int(time.time())}_{idx}{ext}"
        destination = UPLOADS / filename
        content = await f.read()
        with destination.open("wb") as buffer:
            buffer.write(content)

        path = f"/uploads/questions/{filename}"
        img = QImage(
            question_id=question_id,
            path=path,
            order_num=start_order + idx,
        )
        db.add(img)
        saved_images.append(path)
        if not question.media_url:
            question.media_url = path

    db.commit()
    return {"ok": True, "count": len(saved_images), "images": saved_images}


@app.post("/api/questions/{question_id}/unlock")
def unlock_question(
    question_id: int,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    target = db.get(Question, question_id)
    if not target:
        raise HTTPException(status_code=404, detail="Question not found")

    act_key = target.activity_key or "find-object"
    act = db.scalar(select(Activity).where(Activity.key == act_key))
    duration = act.duration_seconds if act else 60

    # 1. Close any other currently active questions in this activity
    other_active = db.scalars(
        select(Question).where(
            Question.status == "active",
            Question.id != question_id,
            (Question.activity_key == act_key) | (Question.activity_key.is_(None) if act_key == "find-object" else False)
        )
    ).all()
    for other in other_active:
        other.status = "closed"

    target.status = "active"
    target.started_at = now()
    target.ends_at = target.started_at + timedelta(seconds=duration)

    # 2. Reset buzzers for this question so count starts from 0
    db.execute(
        delete(Buzzer).where(
            (Buzzer.round_key == f"object:{question_id}")
            | (Buzzer.round_key == f"q_{question_id}")
            | (Buzzer.round_key == f"{act_key}:{question_id}")
        )
    )

    db.commit()
    db.refresh(target)

    return {
        "ok": True,
        "question": {
            "id": target.id,
            "number": target.number,
            "activity_key": target.activity_key,
            "status": target.status,
            "started_at": to_iso(target.started_at),
            "ends_at": to_iso(target.ends_at),
            "hint": target.hint,
        },
    }


@app.post("/api/questions/{question_id}/extend")
def extend_question(
    question_id: int,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    target = db.get(Question, question_id)
    if not target:
        raise HTTPException(status_code=404, detail="Question not found")

    base = target.ends_at if (target.ends_at and target.ends_at > now()) else now()
    target.ends_at = base + timedelta(seconds=30)
    target.status = "active"
    db.commit()
    db.refresh(target)

    return {
        "ok": True,
        "question": {
            "id": target.id,
            "number": target.number,
            "activity_key": target.activity_key,
            "status": target.status,
            "started_at": to_iso(target.started_at),
            "ends_at": to_iso(target.ends_at),
            "hint": target.hint,
        },
    }


@app.post("/api/questions/{question_id}/lock")
def lock_question(
    question_id: int,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    target = db.get(Question, question_id)
    if not target:
        raise HTTPException(status_code=404, detail="Question not found")

    target.status = "locked"
    target.started_at = None
    target.ends_at = None
    db.commit()
    return {"ok": True, "question": {"id": target.id, "status": target.status}}


@app.post("/api/questions/{question_id}/close")
def close_question(
    question_id: int,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    target = db.get(Question, question_id)
    if not target:
        raise HTTPException(status_code=404, detail="Question not found")

    target.status = "closed"
    db.commit()
    return {"ok": True, "question": {"id": target.id, "status": target.status}}


@app.delete("/api/questions/{question_id}")
def delete_question(
    question_id: int,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    """Delete a question regardless of status (active, closed, locked, or completed)."""
    target = db.get(Question, question_id)
    if not target:
        raise HTTPException(status_code=404, detail="Question not found")

    act_key = target.activity_key or "find-object"

    # 1. Delete associated image/media files from disk & QImage records
    images = db.scalars(
        select(QImage).where(QImage.question_id == question_id)
    ).all()
    for img in images:
        if img.path:
            file_name = Path(img.path).name
            file_path = UPLOADS / file_name
            if file_path.is_file():
                try:
                    file_path.unlink()
                except Exception:
                    pass
        db.delete(img)

    # 2. Delete associated submissions
    submissions = db.scalars(
        select(Submission).where(Submission.question_id == question_id)
    ).all()
    for sub in submissions:
        db.delete(sub)

    # 3. Delete buzzer records for this question
    buzzers = db.scalars(
        select(Buzzer).where(
            (Buzzer.round_key == f"object:{question_id}")
            | (Buzzer.round_key == f"q_{question_id}")
            | (Buzzer.round_key == f"{act_key}:{question_id}")
            | (Buzzer.round_key == f"object:{target.number}")
            | (Buzzer.round_key == f"{act_key}:{target.number}")
            | (Buzzer.round_key.like(f"%:{question_id}"))
        )
    ).all()
    for b in buzzers:
        db.delete(b)

    # 4. Delete score records for this question
    scores = db.scalars(
        select(Score).where(
            (Score.round_key == f"object:{question_id}")
            | (Score.round_key == f"q_{question_id}")
            | (Score.round_key == f"{act_key}:{question_id}")
            | (Score.round_key == f"object:{target.number}")
            | (Score.round_key == f"{act_key}:{target.number}")
            | (Score.round_key.like(f"%:{question_id}"))
        )
    ).all()
    for sc in scores:
        db.delete(sc)

    # 5. Delete the question itself
    db.delete(target)
    db.commit()

    return {"ok": True, "deleted_id": question_id}


# ============================================================
# SUBMISSIONS
# ============================================================

@app.get("/api/submissions")
def get_submissions(
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    items = db.scalars(
        select(Submission).order_by(Submission.submitted_at.desc())
    ).all()
    output = []

    for s in items:
        team = db.get(Team, s.team_id)
        user = db.get(User, s.user_id)
        q = db.get(Question, s.question_id)
        output.append(
            {
                "id": s.id,
                "question_id": s.question_id,
                "question_number": q.number if q else None,
                "team_id": s.team_id,
                "team_name": team.name if team else "—",
                "user_id": s.user_id,
                "username": user.username if user else "—",
                "answer_text": s.answer_text,
                "status": s.status,
                "submitted_at": to_iso(s.submitted_at),
                "points": s.points,
            }
        )

    return output


@app.post("/api/submissions")
def submit_answer(
    data: AnswerSubmitIn,
    u: User = Depends(current_user),
    db: Session = Depends(db_dep),
):
    team = team_for(db, u.id)
    if not team:
        raise HTTPException(
            status_code=400,
            detail="You must belong to a team to submit an answer",
        )

    q = db.get(Question, data.question_id)
    if not q or q.status != "active":
        raise HTTPException(
            status_code=400,
            detail="Question is not active for answers",
        )

    if q.ends_at and now() >= q.ends_at:
        q.status = "closed"
        db.commit()
        raise HTTPException(
            status_code=400,
            detail="Time expired for this question",
        )

    sub = Submission(
        question_id=q.id,
        team_id=team.id,
        user_id=u.id,
        answer_text=data.answer_text.strip(),
        status="pending",
        points=0,
    )
    db.add(sub)
    db.commit()
    db.refresh(sub)

    return {"ok": True, "id": sub.id}


@app.post("/api/submissions/{submission_id}/evaluate")
def evaluate_submission(
    submission_id: int,
    data: EvaluateIn,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    sub = db.get(Submission, submission_id)
    if not sub:
        raise HTTPException(status_code=404, detail="Submission not found")

    sub.status = data.status
    sub.points = data.points

    # Update or insert score record in Score table for the leaderboard
    round_key = f"object:{sub.question_id}"
    existing_score = db.scalar(
        select(Score).where(
            Score.team_id == sub.team_id,
            Score.round_key == round_key,
            Score.activity_key == "find-object",
        )
    )
    if existing_score:
        existing_score.points = data.points
        existing_score.updated_by = admin_user.id
    else:
        db.add(
            Score(
                team_id=sub.team_id,
                activity_key="find-object",
                round_key=round_key,
                points=data.points,
                updated_by=admin_user.id,
            )
        )

    db.commit()
    return {"ok": True, "status": sub.status, "points": sub.points}


# ============================================================
# BUZZERS
# ============================================================

@app.get("/api/buzzers")
def get_buzzers(
    filter: Optional[str] = None,
    u: User = Depends(current_user),
    db: Session = Depends(db_dep),
):
    query = select(Buzzer).order_by(Buzzer.pressed_at.asc())
    if filter and filter != "all":
        query = query.where(Buzzer.round_key == filter)

    items = db.scalars(query).all()
    output = []

    for b in items:
        team = db.get(Team, b.team_id)
        if team and team.name == "Admin Staff":
            continue
        user = db.get(User, b.user_id)
        output.append(
            {
                "id": b.id,
                "round_key": b.round_key,
                "round_name": b.round_key,
                "activity_key": b.activity_key,
                "team_id": b.team_id,
                "team_name": team.name if team else "—",
                "user_id": b.user_id,
                "username": user.username if user else "—",
                "pressed_at": to_iso(b.pressed_at),
                "position": b.position,
                "is_winner": b.position == 1,
            }
        )

    if u.role != "admin":
        output = output[:3]

    return output


@app.post("/api/buzzers")
def press_buzzer(
    data: BuzzerIn,
    u: User = Depends(current_user),
    db: Session = Depends(db_dep),
):
    if u.role == "admin":
        raise HTTPException(
            status_code=400,
            detail="Admins are organizers and cannot participate as a competition team.",
        )

    team = team_for(db, u.id)
    if not team:
        raise HTTPException(
            status_code=400,
            detail="You must belong to a team to press the buzzer. Go to Teams page to create or join a team.",
        )

    existing = db.scalar(
        select(Buzzer).where(
            Buzzer.round_key == data.round_key,
            Buzzer.team_id == team.id,
        )
    )
    if existing:
        raise HTTPException(
            status_code=409,
            detail="Your team has already pressed the buzzer for this round",
        )

    count = db.scalar(
        select(func.count(Buzzer.id)).where(Buzzer.round_key == data.round_key)
    ) or 0

    position = count + 1
    buzzer = Buzzer(
        round_key=data.round_key,
        activity_key=data.activity_key or "find-object",
        team_id=team.id,
        user_id=u.id,
        position=position,
    )
    db.add(buzzer)
    db.commit()
    db.refresh(buzzer)

    return {
        "ok": True,
        "id": buzzer.id,
        "position": buzzer.position,
        "is_winner": buzzer.position == 1,
    }


@app.post("/api/buzzers/reset")
@app.delete("/api/buzzers")
def reset_buzzers(
    round_key: Optional[str] = None,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    if round_key and round_key != "all":
        db.execute(delete(Buzzer).where(Buzzer.round_key == round_key))
    else:
        db.execute(delete(Buzzer))
    db.commit()
    return {"ok": True, "message": "Buzzer log cleared"}


# ============================================================
# LIGHTNING SPEECH TOPICS
# ============================================================

@app.get("/api/lightning/topics")
def get_topics(
    u: User = Depends(current_user),
    db: Session = Depends(db_dep),
):
    items = db.scalars(select(Topic).order_by(Topic.id.desc())).all()
    output = []
    current_time = now()

    for t in items:
        if t.status == "active" and t.ends_at and current_time >= t.ends_at:
            t.status = "completed"

        user = db.get(User, t.user_id)
        is_owner = (u.id == t.user_id)
        is_admin = (u.role == "admin")

        output.append(
            {
                "id": t.id,
                "user_id": t.user_id,
                "username": user.username if user else "—",
                "text": t.text if (is_admin or is_owner) else "Private topic",
                "status": t.status,
                "started_at": to_iso(t.started_at),
                "ends_at": to_iso(t.ends_at),
                "resolved": t.resolved,
            }
        )

    db.commit()
    return output


@app.post("/api/lightning/topics")
def assign_topic(
    data: TopicIn,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    user = db.get(User, data.user_id)
    if not user or not user.is_active:
        raise HTTPException(status_code=404, detail="User not found")

    act = db.scalar(select(Activity).where(Activity.key == "lightning-speech"))
    duration = act.duration_seconds if act else 120

    for other in db.scalars(select(Topic).where(Topic.status == "active")).all():
        other.status = "completed"

    start_time = now()
    end_time = start_time + timedelta(seconds=duration)

    existing = db.scalar(select(Topic).where(Topic.user_id == data.user_id))
    if existing:
        existing.text = data.text.strip()
        existing.status = "active"
        existing.started_at = start_time
        existing.ends_at = end_time
        topic = existing
    else:
        topic = Topic(
            user_id=data.user_id,
            text=data.text.strip(),
            status="active",
            started_at=start_time,
            ends_at=end_time,
        )
        db.add(topic)

    db.commit()
    db.refresh(topic)

    return {
        "ok": True,
        "id": topic.id,
        "user_id": topic.user_id,
        "text": topic.text,
        "status": topic.status,
        "started_at": to_iso(topic.started_at),
        "ends_at": to_iso(topic.ends_at),
    }


@app.post("/api/lightning/topics/{topic_id}/unlock")
def unlock_topic(
    topic_id: int,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    act = db.scalar(select(Activity).where(Activity.key == "lightning-speech"))
    duration = act.duration_seconds if act else 120

    for other in db.scalars(select(Topic).where(Topic.status == "active")).all():
        other.status = "completed"

    topic = db.get(Topic, topic_id)
    if not topic:
        raise HTTPException(status_code=404, detail="Topic not found")

    topic.status = "active"
    topic.started_at = now()
    topic.ends_at = topic.started_at + timedelta(seconds=duration)
    db.commit()
    db.refresh(topic)

    return {
        "ok": True,
        "topic": {
            "id": topic.id,
            "status": topic.status,
            "ends_at": to_iso(topic.ends_at),
        },
    }


@app.post("/api/lightning/topics/{topic_id}/lock")
def lock_topic(
    topic_id: int,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    topic = db.get(Topic, topic_id)
    if not topic:
        raise HTTPException(status_code=404, detail="Topic not found")

    topic.status = "locked"
    topic.started_at = None
    topic.ends_at = None
    db.commit()
    db.refresh(topic)
    return {"ok": True, "topic": {"id": topic.id, "status": topic.status}}


@app.post("/api/lightning/topics/{topic_id}/extend")
def extend_topic(
    topic_id: int,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    topic = db.get(Topic, topic_id)
    if not topic:
        raise HTTPException(status_code=404, detail="Topic not found")

    base = topic.ends_at if (topic.ends_at and topic.ends_at > now()) else now()
    topic.ends_at = base + timedelta(seconds=30)
    topic.status = "active"
    db.commit()
    db.refresh(topic)

    return {
        "ok": True,
        "topic": {
            "id": topic.id,
            "status": topic.status,
            "ends_at": to_iso(topic.ends_at),
        },
    }


@app.post("/api/lightning/topics/{topic_id}/close")
def close_topic(
    topic_id: int,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    topic = db.get(Topic, topic_id)
    if not topic:
        raise HTTPException(status_code=404, detail="Topic not found")

    topic.status = "completed"
    db.commit()
    return {"ok": True, "topic": {"id": topic.id, "status": topic.status}}


@app.delete("/api/lightning/topics/{topic_id}")
def delete_topic(
    topic_id: int,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    """Delete a lightning speech topic regardless of status."""
    topic = db.get(Topic, topic_id)
    if not topic:
        raise HTTPException(status_code=404, detail="Topic not found")

    # Delete related scores
    db.query(Score).filter(
        Score.round_key == f"speech:{topic_id}"
    ).delete(synchronize_session=False)

    # Delete related buzzers
    db.query(Buzzer).filter(
        Buzzer.round_key == f"speech:{topic_id}"
    ).delete(synchronize_session=False)

    db.delete(topic)
    db.commit()

    return {"ok": True, "deleted_id": topic_id}


# ============================================================
# LEADERBOARD, SCORES & EXCEL IMPORT
# ============================================================

@app.get("/api/leaderboard")
def get_leaderboard(
    u: User = Depends(current_user),
    db: Session = Depends(db_dep),
):
    teams = db.scalars(
        select(Team).where(Team.name != "Admin Staff").order_by(Team.id)
    ).all()
    rows = []

    all_scores = db.scalars(select(Score)).all()
    scores_by_team = {}
    for s in all_scores:
        scores_by_team.setdefault(s.team_id, []).append({
            "id": s.id,
            "round_key": s.round_key,
            "activity_key": s.activity_key,
            "points": s.points,
        })

    for team in teams:
        manual_score = db.scalar(
            select(func.sum(Score.points)).where(Score.team_id == team.id)
        ) or 0

        imported = db.scalar(
            select(ImportedResult).where(
                func.lower(ImportedResult.team_name) == func.lower(team.name)
            )
        )
        imported_score = imported.total if imported else 0

        total_points = manual_score + imported_score
        rows.append(
            {
                "team_id": team.id,
                "team_name": team.name,
                "name": team.name,
                "points": total_points,
                "score": total_points,
                "scores": scores_by_team.get(team.id, []),
            }
        )

    rows.sort(key=lambda r: r["points"], reverse=True)

    for idx, row in enumerate(rows, 1):
        row["rank"] = idx

    return {"rows": rows}


@app.get("/api/scores")
def get_scores(
    round_key: Optional[str] = None,
    activity_key: Optional[str] = None,
    u: User = Depends(current_user),
    db: Session = Depends(db_dep),
):
    query = select(Score).order_by(Score.id.desc())
    if round_key and round_key != "all":
        query = query.where(Score.round_key == round_key)
    if activity_key and activity_key != "all":
        query = query.where(Score.activity_key == activity_key)

    items = db.scalars(query).all()
    output = []
    for s in items:
        team = db.get(Team, s.team_id)
        user = db.get(User, s.updated_by)
        output.append(
            {
                "id": s.id,
                "team_id": s.team_id,
                "team_name": team.name if team else "—",
                "activity_key": s.activity_key,
                "round_key": s.round_key,
                "points": s.points,
                "updated_by": user.username if user else "—",
            }
        )
    return output


@app.post("/api/scores")
def add_score(
    data: ScoreIn,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    team = db.get(Team, data.team_id)
    if not team:
        raise HTTPException(status_code=404, detail="Team not found")

    act_key = data.activity_key or "find-object"
    existing = db.scalar(
        select(Score).where(
            Score.team_id == data.team_id,
            Score.round_key == data.round_key,
            Score.activity_key == act_key,
        )
    )
    if existing:
        existing.points = data.points
        existing.updated_by = admin_user.id
        db.commit()
        db.refresh(existing)
        return {"ok": True, "id": existing.id, "points": existing.points, "updated": True}
    else:
        score = Score(
            team_id=data.team_id,
            activity_key=act_key,
            round_key=data.round_key,
            points=data.points,
            updated_by=admin_user.id,
        )
        db.add(score)
        db.commit()
        db.refresh(score)
        return {"ok": True, "id": score.id, "points": score.points, "created": True}


@app.delete("/api/scores/{score_id}")
def delete_score(
    score_id: int,
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    score = db.get(Score, score_id)
    if not score:
        raise HTTPException(status_code=404, detail="Score record not found")
    db.delete(score)
    db.commit()
    return {"ok": True, "deleted_id": score_id}


@app.post("/api/results/import")
@app.post("/api/leaderboard/import")
async def import_results(
    file: UploadFile = File(...),
    admin_user: User = Depends(admin),
    db: Session = Depends(db_dep),
):
    try:
        wb = load_workbook(filename=file.file)
        sheet = wb.active
        headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        header_map = {
            str(h).strip().lower(): idx
            for idx, h in enumerate(headers)
            if h is not None
        }

        team_idx = header_map.get("team")
        obj_idx = header_map.get("find the object")
        speech_idx = header_map.get("lightning speech")
        total_idx = header_map.get("total")

        if team_idx is None:
            raise HTTPException(
                status_code=400,
                detail="Excel file must contain a 'Team' column in the first row",
            )

        count = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[team_idx]:
                continue

            team_name = str(row[team_idx]).strip()
            obj_score = (
                int(row[obj_idx])
                if (obj_idx is not None and row[obj_idx] is not None)
                else 0
            )
            speech_score = (
                int(row[speech_idx])
                if (speech_idx is not None and row[speech_idx] is not None)
                else 0
            )
            total = (
                int(row[total_idx])
                if (total_idx is not None and row[total_idx] is not None)
                else (obj_score + speech_score)
            )

            existing = db.scalar(
                select(ImportedResult).where(
                    func.lower(ImportedResult.team_name) == func.lower(team_name)
                )
            )
            if existing:
                existing.object_score = obj_score
                existing.speech_score = speech_score
                existing.total = total
            else:
                db.add(
                    ImportedResult(
                        team_name=team_name,
                        object_score=obj_score,
                        speech_score=speech_score,
                        total=total,
                    )
                )
            count += 1

        db.commit()
        return {"ok": True, "imported": count}
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=400,
            detail=f"Failed to import Excel: {str(e)}",
        )